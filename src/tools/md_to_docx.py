#!/usr/bin/env python3
"""将 docs/讲稿_*.md 转为 Word。"""

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor


def strip_latex_artifacts(text: str) -> str:
    """去掉残留的 LaTeX 定界符，避免 Word 显示乱码。"""
    text = text.replace(r"\(", "").replace(r"\)", "")
    text = re.sub(r"\\\[[\s\S]*?\\\]", "", text)
    text = text.replace(r"\[", "").replace(r"\]", "")
    text = re.sub(r"\\[a-zA-Z]+", "", text)
    text = text.replace(r"\_", "_").replace(r"\^", "^")
    return text


def add_runs_with_bold(paragraph, text: str) -> None:
    text = strip_latex_artifacts(text)
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part)


def md_to_docx(md_path: Path, out_path: Path) -> None:
    text = md_path.read_text(encoding="utf-8")
    doc = Document()

    for line in text.splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
        elif line.strip() == "---":
            continue
        elif line.startswith("|") and "---" in line:
            continue
        elif line.startswith("|"):
            p = doc.add_paragraph(line.strip())
            p.paragraph_format.space_after = Pt(2)
        elif line.startswith("- "):
            body = line[2:].strip()
            p = doc.add_paragraph(style="List Bullet")
            add_runs_with_bold(p, body)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            r = p.add_run(line[2:].strip())
            r.italic = True
        elif "【可选" in line or "【备答" in line:
            p = doc.add_paragraph()
            add_runs_with_bold(p, line)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif "【简版" in line:
            p = doc.add_paragraph()
            add_runs_with_bold(p, line)
            for run in p.runs:
                run.bold = True
        elif line.strip().startswith("**") and line.strip().endswith("**") and line.count("**") == 2:
            p = doc.add_paragraph()
            p.add_run(line.strip().strip("*")).bold = True
        elif line.strip():
            p = doc.add_paragraph()
            add_runs_with_bold(p, line)
            p.paragraph_format.space_after = Pt(4)

    doc.save(str(out_path))
    print(f"已生成: {out_path}")


def main() -> None:
    import sys

    from paths import DOC_SCRIPT_FULL, DOCS, ensure_dirs

    ensure_dirs()
    if len(sys.argv) >= 2:
        stem = sys.argv[1]
        md = DOCS / f"{stem}.md"
        out = DOCS / f"{stem}.docx"
    else:
        md = DOC_SCRIPT_FULL
        out = DOC_SCRIPT_FULL.with_suffix(".docx")
    if not md.is_file():
        raise SystemExit(f"找不到: {md}")
    md_to_docx(md, out)


if __name__ == "__main__":
    import _bootstrap  # noqa: F401

    main()
