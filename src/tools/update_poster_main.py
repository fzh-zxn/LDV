#!/usr/bin/env python3
"""重写 LDV_海报排版与文字.docx：顶栏 + Introduction + Main Approach。"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def delete_paragraph(paragraph) -> None:
    el = paragraph._element
    el.getparent().remove(el)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_body(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    p.paragraph_format.space_after = Pt(5)


def add_bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def find_para_index(doc: Document, start: int, needle: str) -> int:
    for i in range(start, len(doc.paragraphs)):
        if needle in doc.paragraphs[i].text:
            return i
    return -1


def main() -> None:
    from paths import DOC_POSTER_TEXT

    path = DOC_POSTER_TEXT
    if not path.is_file():
        raise FileNotFoundError(path)

    doc = Document(str(path))

    # 锚点：从「中文标题」到「Data Processing」之前删掉，保留其后（数据处理/结论/图等）
    i_title = find_para_index(doc, 0, "中文标题")
    if i_title < 0:
        i_title = find_para_index(doc, 0, "激光多普勒")
    i_data = find_para_index(doc, 0, "Data Processing")
    if i_data < 0:
        i_data = find_para_index(doc, 0, "数据处理")

    if i_title < 0 or i_data < 0:
        raise RuntimeError(f"找不到锚点: title={i_title}, data={i_data}")

    # 倒序删除，避免图片段落在范围内被删掉——只删纯文字段
    to_del = []
    for i in range(i_title, i_data):
        p = doc.paragraphs[i]
        if "w:drawing" not in p._element.xml and "pic:pic" not in p._element.xml:
            to_del.append(p)
    for p in reversed(to_del):
        delete_paragraph(p)

    # 在「Data Processing」段前插入新内容（重新定位）
    i_data = find_para_index(doc, 0, "Data Processing")
    if i_data < 0:
        i_data = find_para_index(doc, 0, "数据处理")
    anchor = doc.paragraphs[i_data]

    new_blocks: list[tuple[str, str]] = [
        ("h0", "顶栏"),
        ("t", "中文标题：激光多普勒远程侦听"),
        ("t", "英文标题：Remote Acoustic Sensing via Laser Doppler Vibrometry"),
        ("t", "副标题（可选）：正交鉴频 · Fs 标定 · 音阶/人声双验证"),
        ("t", "小组成员：（填写）"),
        ("t", "指导教师：（学院　姓名）"),
        ("blank", ""),
        ("h1", "Introduction　实验背景"),
        (
            "p",
            "振动是工程与自然界的常见物理量；精确测振对安防监听、结构监测等很重要。"
            "接触式传感器需贴附目标，且易受电磁干扰，不适用于远距离、非合作目标。",
        ),
        (
            "p",
            "激光自混合（回馈）干涉可对微小位移高度敏感：目标散射光回馈激光腔，"
            "与腔内光场干涉，光电探测器（PD）输出随位移变化。"
            "光路中加入声光移频器（AOM）产生稳定载波，可减轻低频漂移、方向模糊与灵敏度衰落。"
            "锁相放大器对 PD 信号做正交解调，得到同相/正交分量 I(t)、Q(t)——"
            "这是后续「声音/振动还原」的输入（算法见 Main Approach）。",
        ),
        ("note", "【插图】图1　实验光路示意图（AOM + 自混合干涉）"),
        ("blank", ""),
        ("h1", "Main Approach　核心方法"),
        ("h2", "1. 传统方法的局限"),
        (
            "p",
            "常见做法：对 I/Q 用 atan2 求相位，再相位解包裹并微分得到速度。"
            "在散斑噪声下相位易在 ±π 附近突变，解调结果出现全频带伪峰，"
            "微弱音阶或人声常被噪声淹没。",
        ),
        ("h2", "2. 本工作思路"),
        (
            "p",
            "不对 I/Q 求相位，而在笛卡尔坐标下对去均值后的 I、Q 做差分交叉相乘，"
            "直接得到与靶面角速度（振动速度）成正比的量，从原理上规避相位跳变。",
        ),
        ("h2", "3. 正交鉴频公式"),
        ("p", "v(t) ∝ ( I·dQ − Q·dI ) / ( I² + Q² + ε )"),
        ("p", "ε 为极小正数，防止光强趋零时分母发散。"),
        ("h2", "4. 方法特点"),
        ("b", "① 稳定：无需相位解包裹，适合散斑环境。"),
        ("b", "② 可验证：音阶 500 / 1 k / 2 kHz 标定与连续人声并列检验同一 DSP 链路。"),
        ("b", "③ 可扩展：拼接淡化、Fs 标定、分级降噪可模块化替换。"),
        (
            "note",
            "【插图】图2　建议：左「atan + 解包裹」相位突变示意；右「交叉相乘鉴频」框图（本工作）",
        ),
        ("blank", ""),
    ]

    body = anchor._element.getparent()
    insert_pos = list(body).index(anchor._element)

    tmp = Document()
    for kind, text in new_blocks:
        if kind == "blank":
            tmp.add_paragraph()
        elif kind == "h0":
            p = tmp.add_paragraph()
            p.add_run(text).bold = True
        elif kind == "h1":
            tmp.add_heading(text, level=1)
        elif kind == "h2":
            tmp.add_heading(text, level=2)
        elif kind == "t":
            p = tmp.add_paragraph()
            r = p.add_run(text)
            if text.startswith("中文标题"):
                r.bold = True
                r.font.size = Pt(16)
        elif kind == "note":
            p = tmp.add_paragraph()
            p.add_run(text).italic = True
        elif kind == "b":
            tmp.add_paragraph(text, style="List Bullet")
        else:
            tmp.add_paragraph(text)

    for para in reversed(tmp.paragraphs):
        body.insert(insert_pos, para._element)

    doc.save(str(path))
    print(f"已更新: {path}")
    print(f"重写范围: 顶栏 → Introduction → Main Approach（保留 Data Processing 及之后）")


if __name__ == "__main__":
    import _bootstrap  # noqa: F401

    main()
