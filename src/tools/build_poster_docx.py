#!/usr/bin/env python3
"""生成 LDV 海报排版与文字 Word 文档。"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    p.paragraph_format.space_after = Pt(6)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def main() -> None:
    from paths import DOC_POSTER_TEXT, ensure_dirs

    ensure_dirs()
    out = DOC_POSTER_TEXT

    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)

    # 封面说明
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("激光多普勒远程侦听 — 学术海报")
    r.bold = True
    r.font.size = Pt(22)

    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = t2.add_run("排版结构 · 分栏文字 · 图件清单 · 讲稿提纲")
    r2.font.size = Pt(14)

    doc.add_paragraph(
        "说明：版式参照《基于存算一体架构的深度卷积神经网络实现_海报.pptx》"
        "（Introduction → Main Approach → Data Processing → Results → Outlook）。"
        "可直接将各栏文字复制到 PPT；图件见本文件夹内 PNG。"
    )

    add_heading(doc, "使用说明", 1)
    add_bullets(
        doc,
        [
            "海报建议单页竖版 A0，三主栏 + 顶栏 + 底栏 Outlook。",
            "结果图共 3 张数据处理图 + 图1光路 + 图2公式 + 图3流程（后三张可手绘框图）。",
            "音阶：ldv_triptych_spec.png、ldv_triptych_welch.png；人声：ldv_voice_diptych_spec.png。",
            "小组成员、指导教师、单位请在顶栏自行填写。",
        ],
    )

    # 顶栏
    add_heading(doc, "一、顶栏（Poster Header）", 1)
    add_para(doc, "中文标题：激光多普勒远程侦听", bold=True)
    add_para(doc, "英文标题：Remote Acoustic Sensing via Laser Doppler Vibrometry")
    add_para(doc, "小组成员：（填写）")
    add_para(doc, "指导教师：（学院　姓名）")
    add_para(doc, "副标题/卖点（可选）：正交鉴频 · Fs 标定 · 音阶/人声双验证")

    # Introduction
    add_heading(doc, "二、Introduction  实验背景（左栏）", 1)
    add_para(
        doc,
        "振动是自然界与工程中最常见的物理现象之一；精确测振对结构健康监测、"
        "安防与微声学检测均具有重要意义。接触式传感器需贴附目标且易受电磁干扰，"
        "难以用于远距离、非合作目标。",
    )
    add_para(
        doc,
        "光学非接触测量精度高、无侵入。激光自混合/回馈干涉将目标微小位移映射为 "
        "PD 电流波动；在回馈光路中插入声光移频器（AOM）引入稳定载波 fm，"
        "可缓解低频漂移、方向模糊与灵敏度衰落。锁相放大器输出 I/Q 正交分量，"
        "为后续振动与声音还原提供基础。",
    )
    add_para(doc, "【插图】图1  激光自混合干涉实验光路示意图（自备照片或示意图）", bold=True)

    # Main Approach
    add_heading(doc, "三、Main Approach  核心方法（中栏上部）", 1)
    add_para(doc, "传统困境", bold=True)
    add_para(
        doc,
        "在散斑噪声下，atan2(Q,I) 相位解包裹易发散，微弱振动被全频带伪峰与白噪声淹没。",
    )
    add_para(doc, "正交鉴频（本工作）", bold=True)
    add_para(doc, "去均值后，由差分交叉相乘直接估计角速度（正比于靶面振动速度）：")
    add_para(doc, "v(t) ∝ (I·dQ − Q·dI) / (I² + Q² + ε)")
    add_para(doc, "Main Pros", bold=True)
    add_bullets(
        doc,
        [
            "规避相位跳变：笛卡尔域鉴频，无需相位解包裹。",
            "双目标验证：音阶标定（500/1/2 kHz）与连续人声并列验证链路。",
            "模块化 DSP：拼接淡化、Fs 标定、分级降噪，便于扩展。",
        ],
    )
    add_para(doc, "【插图】图2  正交鉴频公式与传统 atan 解调对比（小图）", bold=True)

    # Data Processing
    add_heading(doc, "四、Data Processing  数据处理（中栏下部）", 1)
    add_para(doc, "音阶信号 — 三步（stream_00000~00002.mat 拼接）", bold=True)

    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    hdr = ["阶段", "名称", "内容"]
    for i, h in enumerate(hdr):
        table.rows[0].cells[i].text = h
    rows = [
        ("Stage 1", "信号还原", "I/Q 拼接 + 50 ms 接缝淡化 → 鉴频 → 98% 削峰 → 宽带通 → Fs 标定"),
        ("Stage 2", "简单降噪", "前 0.5 s 谱减 + 固定带通 300–2500 Hz"),
        ("Stage 3", "最终降噪", "谱减 + 分峰自适应窄带（按段检测 500/1k/2k Hz）"),
    ]
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            table.rows[r].cells[c].text = val

    doc.add_paragraph()
    add_para(doc, "人声信号 — 两步（ldv_iq_demodulated.wav）", bold=True)
    table2 = doc.add_table(rows=3, cols=2)
    table2.style = "Table Grid"
    table2.rows[0].cells[0].text = "步骤"
    table2.rows[0].cells[1].text = "内容"
    table2.rows[1].cells[0].text = "①"
    table2.rows[1].cells[1].text = "原始鉴频波形（未降噪）"
    table2.rows[2].cells[0].text = "②"
    table2.rows[2].cells[1].text = (
        "50–2000 Hz 带通 + 全段 20% 分位噪声底谱减（不用开头 0.5 s 作噪声参考）"
    )

    doc.add_paragraph()
    add_para(
        doc,
        "Fs 标定说明：三段 I/Q 点数不同，不可用「总点数/30 s」；"
        "以中间段主峰对齐 1000 Hz，使谱线约为 500 / 1000 / 2000 Hz。",
    )
    add_para(doc, "【插图】图3  音阶三步 / 人声两步 流程框图（建议 PPT 自绘）", bold=True)

    # Results
    add_heading(doc, "五、Results and Discussion  结果与讨论（右栏）", 1)
    add_para(doc, "【插图】图4  ldv_triptych_spec.png", bold=True)
    add_para(doc, "音阶三级处理语谱对比；竖虚线为三份 MAT 分界。")
    add_para(doc, "【插图】图5  ldv_triptych_welch.png", bold=True)
    add_para(doc, "音阶 Welch 功率谱；标定后主峰约 500 Hz / 1000 Hz / 2008 Hz。")
    add_para(doc, "【插图】图6  ldv_voice_diptych_spec.png", bold=True)
    add_para(doc, "人声降噪前后语谱（0–2000 Hz）；展示连续语音可听化与降噪效果。")

    add_para(doc, "结论（三条，仿参考海报编号）", bold=True)
    add_bullets(
        doc,
        [
            "成功搭建 AOM–自混合光路，获取 IQ 信号并导出 WAV。",
            "音阶主峰与标称频率一致，Fs 标定策略有效。",
            "人声两步降噪后语谱结构更清晰，可现场耳机演示。",
        ],
    )

    # Outlook
    add_heading(doc, "六、Outlook  展望（底栏通栏）", 1)
    add_bullets(
        doc,
        [
            "算法：深度学习去噪与经典 DSP 融合，适应极端信噪比。",
            "应用：无损检测、MEMS 模态分析、复合材料缺陷诊断、安防远距离拾音。",
            "系统：更高采样率、实时流式处理与多目标跟踪。",
        ],
    )

    # 图注
    add_heading(doc, "七、图注汇总（粘贴到图下方）", 1)
    captions = [
        "图1  激光自混合干涉实验光路示意图",
        "图2  正交鉴频原理与传统相位解调对比",
        "图3  音阶（三步）与人声（两步）数据处理流程",
        "图4  音阶信号三级处理语谱图对比",
        "图5  音阶信号三级处理 Welch 功率谱对比",
        "图6  人声信号降噪前后语谱图对比（0–2000 Hz）",
    ]
    for c in captions:
        doc.add_paragraph(c)

    # 讲稿
    add_heading(doc, "八、答辩讲稿提纲（约 6 分钟）", 1)
    script = [
        ("0:00–0:40  开场", "题目、非接触 LDV、音阶+人声双验证。"),
        ("0:40–2:00  Introduction", "接触/非接触对比；自混合+AOM；锁相 IQ。光路图。"),
        ("2:00–3:30  Main Approach", "atan 失效；鉴频公式；三条 Main Pros。"),
        ("3:30–5:00  Data Processing + Results", "音阶三步、人声两步；对着图4–6讲。"),
        ("5:00–5:30  结论", "光路、频率标定、人声可听化。"),
        ("5:30–6:00  Outlook", "AI、应用拓展。"),
    ]
    for title, body in script:
        p = doc.add_paragraph()
        p.add_run(title + "：").bold = True
        p.add_run(body)

    add_heading(doc, "九、常见问题备答", 1)
    faq = [
        ("为何不用 atan？", "散斑导致相位突变；交叉相乘在笛卡尔域更稳定。"),
        ("为何不是 30 s？", "三段采样率不一致，需主峰标定 Fs，不能用总点数/30。"),
        ("音阶简单 vs 最终？", "简单=固定带通；最终=按段多峰自适应窄带。"),
        ("人声为何不用 0.5 s 谱减？", "开头可能已是语音；改用全段低分位估计噪声谱。"),
    ]
    for q, a in faq:
        p = doc.add_paragraph()
        p.add_run("Q：" + q).bold = True
        doc.add_paragraph("A：" + a)

    add_heading(doc, "十、本地文件清单", 1)
    files = [
        "ldv_triptych_spec.png / ldv_triptych_welch.png — 音阶三联图",
        "ldv_voice_diptych_spec.png — 人声二联语谱",
        "ldv_30s_scale_test.wav — 音阶还原",
        "ldv_denoised_simple.wav / ldv_denoised.wav — 音阶降噪",
        "ldv_iq_demodulated(1).wav / ldv_voice_denoised.wav — 人声",
        "python run_pipeline.py — 音阶一键处理",
        "python ldv_denoise_voice.py — 人声处理",
    ]
    add_bullets(doc, files)

    doc.save(out)
    print(f"已生成: {out}")


if __name__ == "__main__":
    import _bootstrap  # noqa: F401

    main()
