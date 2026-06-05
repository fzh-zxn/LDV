#!/usr/bin/env python3
"""重写 Word 中「结论」与「Outlook 展望」。"""

from pathlib import Path

from docx import Document
from docx.shared import Pt


def delete_paragraph(paragraph) -> None:
    el = paragraph._element
    el.getparent().remove(el)


def find_para(doc: Document, needle: str, start: int = 0) -> int:
    for i in range(start, len(doc.paragraphs)):
        if needle in doc.paragraphs[i].text:
            return i
    return -1


def has_media(paragraph) -> bool:
    xml = paragraph._element.xml
    return "w:drawing" in xml or "pic:pic" in xml


def main() -> None:
    from paths import DOC_POSTER_TEXT

    path = DOC_POSTER_TEXT
    doc = Document(str(path))

    i_conc = find_para(doc, "结论")
    if i_conc < 0:
        i_conc = find_para(doc, "Results")
    if i_conc < 0:
        raise RuntimeError("未找到「结论」或 Results 锚点，请保留原文标题之一")

    # 删除从「结论」到文末的纯文字段（保留带图片的段）
    to_del = []
    for i in range(i_conc, len(doc.paragraphs)):
        p = doc.paragraphs[i]
        if not has_media(p):
            to_del.append(p)
    for p in reversed(to_del):
        delete_paragraph(p)

    new_blocks = [
        ("h1", "Results and Discussion　结果与讨论"),
        ("h2", "实验结论"),
        (
            "b",
            "① 系统验证：完成 AOM–自混合干涉光路搭建，由锁相放大器获取 IQ 信号，"
            "经正交鉴频与 DSP 后导出可听 WAV，证明非接触 LDV 拾音链路可行。",
        ),
        (
            "b",
            "② 音阶标定：三段 MAT 拼接后经 Fs 标定，语谱与 Welch 主峰约 "
            "500 Hz / 1 kHz / 2 kHz，与标称音阶一致，说明采样率标定策略有效。",
        ),
        (
            "b",
            "③ 人声还原：对鉴频波形做 50–2000 Hz 带通与全段噪声底谱减，"
            "降噪后语谱结构更清晰，可现场耳机试听对比。",
        ),
        (
            "p",
            "分级处理表明：简单降噪可快速抑制宽带噪声；"
            "分峰自适应窄带更适合分段音阶；"
            "人声与音阶并列验证同一鉴频框架在不同信号形态下的适用性。",
        ),
        ("blank", ""),
        ("h1", "Outlook　展望"),
        (
            "b",
            "算法：在经典 DSP（谱减、带通、自适应窄带）基础上引入深度学习去噪，"
            "提升极端信噪比下的可懂度。",
        ),
        (
            "b",
            "应用：拓展至结构无损检测、MEMS 模态分析、复合材料内部缺陷诊断"
            "及安防场景远距离拾音。",
        ),
        (
            "b",
            "系统：提高采集带宽与实时处理能力，优化多段拼接与长期稳定测量。",
        ),
    ]

    tmp = Document()
    for kind, text in new_blocks:
        if kind == "blank":
            tmp.add_paragraph()
        elif kind == "h1":
            tmp.add_heading(text, level=1)
        elif kind == "h2":
            tmp.add_heading(text, level=2)
        elif kind == "b":
            tmp.add_paragraph(text, style="List Bullet")
        else:
            p = tmp.add_paragraph()
            p.add_run(text)
            p.paragraph_format.space_after = Pt(5)

    body = doc.element.body
    for para in reversed(tmp.paragraphs):
        body.append(para._element)

    try:
        doc.save(str(path))
        print(f"已更新: {path}")
    except PermissionError:
        alt = path.with_name(path.stem + "_已改结论展望.docx")
        doc.save(str(alt))
        print(f"原文件被占用，已另存为: {alt}")
        print("请关闭 Word 后可将新文件覆盖原文件，或手动复制内容。")


if __name__ == "__main__":
    import _bootstrap  # noqa: F401

    main()
